"""Acceptance checks for the presentation2 renderer pipeline (build.py,
validate_sections.py, mathimg.py, render_pptx.js, render_html.py).

This runs the build against an isolated temp copy of the schema-conformance
sample section only (presentation2/sections/00_sample.json), never against
the live presentation2/sections/ directory -- that directory is being
written concurrently by other section-content workers, and this check must
be deterministic regardless of the state of their in-progress files.

What this confirms:
  * `python presentation2/build.py --include-sample --sections-dir <isolated>`
    exits 0 (validate -> figures -> equations -> pptx -> html -> office
    validate, per presentation2/SCHEMA.md's renderer contract);
  * the produced pptx opens with python-pptx, its slide count equals the
    sample section's slide count, every slide's rendered text contains its
    JSON `title`, every slide's notes slide matches its JSON `notes`
    verbatim, and every `equation`-layout slide contains at least as many
    picture shapes as it has `equations`;
  * out/presentation2/index.html has the same number of `<section` elements
    as slides, embeds no external `<script src=`, and stays under 16 MB.

Sandbox note: this creates a tempfile.TemporaryDirectory(); inside a
read-only or workspace-write sandbox that can raise PermissionError. If that
is the ONLY failure, report `TESTS: pass (verify_presentation2 skipped:
sandbox temp dir)` per CLAUDE.md and let the orchestrator re-run outside
the sandbox.

CLI: python verify/verify_presentation2.py
Exit code 0 iff every check passes; prints 'N/N presentation2 checks passed'.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
PRESENTATION2 = ROOT / "presentation2"
SAMPLE_JSON = PRESENTATION2 / "sections" / "00_sample.json"
OUT_DIR = ROOT / "out" / "presentation2"
PPTX_PATH = OUT_DIR / "qd_physics_2h.pptx"
HTML_PATH = OUT_DIR / "index.html"

MAX_HTML_BYTES = 16 * 1024 * 1024

checks: list[bool] = []


def ck(ok: bool, name: str) -> None:
    checks.append(bool(ok))
    print(("PASS" if ok else "FAIL"), name)


def main() -> int:
    sample = json.loads(SAMPLE_JSON.read_text(encoding="utf-8"))
    slides_by_id = {s["id"]: s for s in sample["slides"]}
    expected_n_slides = len(sample["slides"])

    with tempfile.TemporaryDirectory() as tmp:
        sections_dir = Path(tmp)
        shutil.copy(SAMPLE_JSON, sections_dir / SAMPLE_JSON.name)

        result = subprocess.run(
            [sys.executable, str(PRESENTATION2 / "build.py"),
             "--include-sample", "--sections-dir", str(sections_dir)],
            cwd=str(ROOT),
        )
        ck(result.returncode == 0, "presentation2/build.py --include-sample exits 0")
        if result.returncode != 0:
            print(f"{sum(checks)}/{len(checks)} presentation2 checks passed")
            return 1

    prs = Presentation(str(PPTX_PATH))
    slides = list(prs.slides)
    ck(len(slides) == expected_n_slides,
       f"pptx slide count == sample slide count ({len(slides)} vs {expected_n_slides})")

    for i, pptx_slide in enumerate(slides):
        json_slide = sample["slides"][i] if i < len(sample["slides"]) else None
        label = json_slide["id"] if json_slide else f"index {i}"

        slide_text = "\n".join(
            shape.text_frame.text for shape in pptx_slide.shapes if shape.has_text_frame
        )
        title_ok = bool(json_slide) and json_slide["title"] in slide_text
        ck(title_ok, f"slide {label}: rendered text contains its JSON title")

        notes_text = ""
        if pptx_slide.has_notes_slide:
            notes_text = pptx_slide.notes_slide.notes_text_frame.text
        notes_ok = bool(json_slide) and notes_text == json_slide["notes"]
        ck(notes_ok, f"slide {label}: notes slide matches its JSON notes verbatim")

        if json_slide and json_slide["layout"] == "equation":
            n_pictures = sum(
                1 for shape in pptx_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            )
            n_equations = len(json_slide.get("equations", []))
            ck(n_pictures >= n_equations,
               f"slide {label}: {n_pictures} pictures >= {n_equations} equations")

    html_text = HTML_PATH.read_text(encoding="utf-8")
    n_sections = len(re.findall(r"<section", html_text))
    ck(n_sections == expected_n_slides,
       f"index.html has the same number of <section> elements as slides "
       f"({n_sections} vs {expected_n_slides})")
    ck("<script src=" not in html_text, "index.html has no external <script src=...>")
    ck(not re.search(r"""(?:src|href)\s*=\s*["']https?://""", html_text),
       "index.html has no external http(s) src/href reference")
    html_bytes = HTML_PATH.stat().st_size
    ck(html_bytes < MAX_HTML_BYTES, f"index.html is under 16 MB (got {html_bytes} bytes)")

    print(f"{sum(checks)}/{len(checks)} presentation2 checks passed")
    return 0 if all(checks) else 1


if __name__ == "__main__":
    try:
        sys.exit(main())
    except PermissionError as e:
        print(f"SKIPPED (sandbox temp dir): {e}")
        print("See CLAUDE.md 'Sandbox note for workers' -- re-run outside the sandbox.")
        sys.exit(0)
